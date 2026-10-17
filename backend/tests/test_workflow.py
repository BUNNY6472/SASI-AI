import asyncio
import time

import pytest
from fastapi.testclient import TestClient

from app.config import Settings
from app.gateway.base import ModelCall, ModelResult, ProviderError
from app.main import create_app
from app.workflow.controller import Controller
from app.workflow.prompts import parse_json


def make_client(**overrides) -> TestClient:
    opts = {"allow_simulated": True, "simulated_delay": 0.0, **overrides}
    settings = Settings(_env_file=None, static_dir="/nonexistent", **opts)
    return TestClient(create_app(settings))


def wait_for(client: TestClient, run_id: str, statuses: set[str], timeout: float = 10) -> dict:
    end = time.time() + timeout
    while time.time() < end:
        run = client.get(f"/api/runs/{run_id}").json()
        if run["status"] in statuses:
            return run
        time.sleep(0.05)
    raise AssertionError(f"run stuck in {run['status']}")


def test_parse_json_handles_fences_and_prose():
    assert parse_json('Sure!\n```json\n{"a": {"b": "}"}}\n```') == {"a": {"b": "}"}}
    assert parse_json("no json here") is None


def test_full_run_in_demo_mode_with_plan_approval():
    with make_client() as client:
        providers = client.get("/api/providers").json()
        assert providers["demo_mode"] is True
        run = client.post("/api/runs", json={"request": "Design a todo app API", "profile": "software"}).json()
        run = wait_for(client, run["id"], {"awaiting_approval"})
        assert run["plan"]["tasks"], "plan has tasks"
        assert run["plan_review"]["reviewer"] != run["plan"]["planner"], "plan reviewed by a different model"

        client.post(f"/api/runs/{run['id']}/approve", json={"note": "go"})
        run = wait_for(client, run["id"], {"completed", "failed"})
        assert run["status"] == "completed", run.get("error")
        for task in run["plan"]["tasks"]:
            assert task["status"] == "accepted"
            assert task["reviewer"] != task["implementer"], "independent cross-review"
            assert task["assurance"] == "evidence-reviewed"
        assert run["final"]["output"]
        assert run["final"]["evaluator"] != run["final"]["integrator"]
        assert run["usage"]["calls"] > 0
        assert run["simulated"] is True


def test_budget_limit_stops_run_and_keeps_state():
    with make_client(max_model_calls_per_run=8) as client:
        run = client.post("/api/runs", json={"request": "Write a report", "require_plan_approval": False}).json()
        run = wait_for(client, run["id"], {"budget_stopped", "completed", "failed"})
        assert run["status"] == "budget_stopped"
        assert run["plan"] is not None


def test_access_code_gate():
    with make_client(app_access_code="secret-code") as client:
        assert client.get("/api/runs").status_code == 401
        assert client.post("/api/login", json={"code": "wrong"}).status_code == 401
        assert client.post("/api/login", json={"code": "secret-code"}).status_code == 200
        assert client.get("/api/runs").status_code == 200


def test_dependency_cycles_are_broken():
    from app.models import Task

    tasks = [Task(id="a", title="a", depends_on=["b"]), Task(id="b", title="b", depends_on=["a"])]
    Controller._break_cycles(tasks)
    assert not (tasks[0].depends_on and tasks[1].depends_on)


def test_failing_provider_is_replaced():
    """A provider that errors is skipped and its work is reassigned; the run still completes."""
    with make_client() as client:
        registry = client.app.state.registry

        class Broken:
            provider = "chatgpt"
            display_name = "ChatGPT"
            simulated = True

            async def list_models(self):
                return ["broken"]

            async def resolve_model(self):
                return "broken"

            async def generate(self, call: ModelCall) -> ModelResult:
                raise ProviderError("chatgpt", "boom")

        registry.adapters["chatgpt"] = Broken()
        run = client.post("/api/runs", json={"request": "Plan a launch", "require_plan_approval": False}).json()
        run = wait_for(client, run["id"], {"completed", "failed"})
        assert run["status"] == "completed", run.get("error")
        assert any(e["type"] == "provider_error" for e in run["events"])


def test_real_mode_needs_keys_before_running():
    with make_client(allow_simulated=False) as client:
        p = client.get("/api/providers").json()
        assert p["needs_keys"] is True
        assert all(not x["connected"] for x in p["providers"])
        r = client.post("/api/runs", json={"request": "Hello there"})
        assert r.status_code == 400
        assert "API key" in r.json()["detail"]


def test_keys_saved_from_app_are_never_returned(monkeypatch):
    with make_client(allow_simulated=False) as client:
        registry = client.app.state.registry

        async def fake_test(provider, key):
            if key.endswith("bad-key-0000"):
                raise RuntimeError(f"401 invalid key {key}")
            return "model-x"

        monkeypatch.setattr(registry, "test_key", fake_test)
        secret = "sk-test-1234567890abcdWXYZ"
        r = client.put("/api/keys/claude", json={"key": secret})
        assert r.status_code == 200, r.text
        assert secret not in r.text and r.json()["hint"] == "…WXYZ"

        bad = client.put("/api/keys/gemini", json={"key": "sk-bad-key-0000"})
        assert bad.status_code == 400 and "sk-bad-key-0000" not in bad.text

        assert client.put("/api/keys/chatgpt", json={"key": "short"}).status_code == 400

        for path in ("/api/keys", "/api/providers"):
            assert secret not in client.get(path).text
        provs = {p["provider"]: p for p in client.get("/api/providers").json()["providers"]}
        assert provs["claude"]["connected"] and provs["claude"]["model"] == "model-x"
        assert not provs["gemini"]["connected"]

        assert client.delete("/api/keys/claude").status_code == 200
        assert not any(p["configured"] for p in client.get("/api/keys").json())


def test_files_all_three_analyse_compare_and_delete():
    """Files reach every model, each analyses on its own, they cross-check, and delete removes everything."""
    with make_client() as client:
        csv_body = b"region,sales\nnorth,120\nsouth,80\n"
        up = client.post("/api/files", content=csv_body, headers={"X-File-Name": "sales.csv", "Content-Type": "text/csv"})
        assert up.status_code == 200, up.text
        f = up.json()
        assert f["kind"] == "text"
        img = client.post("/api/files", content=b"\x89PNG\r\n\x1a\nfake", headers={"X-File-Name": "chart.png", "Content-Type": "image/png"}).json()
        empty = client.post("/api/files", content=b"", headers={"X-File-Name": "x.csv", "Content-Type": "text/csv"})
        assert empty.status_code == 400

        run = client.post("/api/runs", json={"request": "Analyse the sales data", "profile": "research",
                                             "require_plan_approval": False, "file_ids": [f["id"], img["id"]]}).json()
        assert [x["name"] for x in run["files"]] == ["sales.csv", "chart.png"]
        run = wait_for(client, run["id"], {"completed", "failed"})
        assert run["status"] == "completed", run.get("error")
        assert {a["provider"] for a in run["analyses"]} == {"claude", "chatgpt", "gemini"}
        assert all("sales.csv" in a["output"] for a in run["analyses"]), "every model saw the files"
        assert len(run["cross_checks"]) == 3
        assert run["live"] == {}, "live text is cleared when the run ends"

        store = client.app.state.store
        stages = {o["stage"] for o in store.outputs.get(run["id"], [])}
        assert {"analysis", "cross_check", "task", "review", "integrate"} <= stages, stages

        got = client.get(f"/api/runs/{run['id']}/files/{f['id']}")
        assert got.status_code == 200 and got.content == csv_body

        # A staged file can only be used once.
        again = client.post("/api/runs", json={"request": "again", "file_ids": [f["id"]]})
        assert again.status_code == 400

        d = client.delete(f"/api/runs/{run['id']}")
        assert d.status_code == 200 and d.json()["deleted"] is True
        assert client.get(f"/api/runs/{run['id']}").status_code == 404
        assert run["id"] not in store.outputs


def test_csv_profile_and_type_checks():
    from app.files import FileRejected, prepare

    f = prepare("data.csv", "", b"a,b\n1,x\n3,y\n")
    assert "Rows: 2" in f.text and "min 1, max 3" in f.text
    with pytest.raises(FileRejected):
        prepare("empty.csv", "", b"")
    exe = prepare("tool.exe", "application/octet-stream", b"MZ\x00\x01\x02")
    assert exe.kind == "other", "any type is accepted and kept"
    assert prepare("clip.mov", "", b"0000").kind == "video"
    assert prepare("logo.svg", "image/svg+xml", b"<svg/>").kind == "text"


def test_data_and_document_files_become_readable_text():
    import io
    import json
    import zipfile

    import openpyxl
    import pyarrow as pa
    import pyarrow.parquet as pq

    from app.files import prepare

    buf = io.BytesIO()
    pq.write_table(pa.table({"city": ["Hyd", "Vizag", "Hyd"], "sales": [10.5, 20.0, 30.0]}), buf)
    f = prepare("sales.parquet", "", buf.getvalue())
    assert f.kind == "text" and "Rows: 3" in f.text and "max 30" in f.text and "Vizag" in f.text

    wb = openpyxl.Workbook()
    wb.active.append(["item", "qty"])
    wb.active.append(["pen", 4])
    xbuf = io.BytesIO()
    wb.save(xbuf)
    f = prepare("stock.xlsx", "", xbuf.getvalue())
    assert "Sheet" in f.text and "pen" in f.text and "max 4" in f.text

    f = prepare("data.json", "application/json", json.dumps([{"a": 1, "b": "x"}, {"a": 5, "b": "y"}]).encode())
    assert "2 items" in f.text and "max 5" in f.text

    dbuf = io.BytesIO()
    with zipfile.ZipFile(dbuf, "w") as z:
        z.writestr("word/document.xml", '<w:document><w:p><w:r><w:t>Hello</w:t></w:r><w:r><w:t xml:space="preserve"> SASI</w:t></w:r></w:p></w:document>')
    assert "Hello SASI" in prepare("note.docx", "", dbuf.getvalue()).text

    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w") as z:
        z.writestr("inner/a.csv", "x,y\n1,2\n")
        z.writestr("inner/b.parquet", buf.getvalue())
    f = prepare("bundle.zip", "", zbuf.getvalue())
    assert "inner/a.csv" in f.text and "Vizag" in f.text


def test_cannot_delete_a_running_run():
    with make_client(simulated_delay=0.3) as client:
        run = client.post("/api/runs", json={"request": "Long task", "require_plan_approval": True}).json()
        assert client.delete(f"/api/runs/{run['id']}").status_code == 409
        client.post(f"/api/runs/{run['id']}/cancel")


def test_bigquery_file_chunks_round_trip(monkeypatch):
    """Large files are split into chunks for BigQuery and joined back in order; delete clears every table."""
    import base64 as b64
    import datetime

    from app.files import StoredFile
    from app.store import bigquery as bqmod

    class FakeDB:
        def __init__(self):
            self.tables = {bqmod.CHECKPOINTS: [], bqmod.OUTPUTS: [], bqmod.FILES: []}

        def ensure(self):
            pass

        def table(self, name):
            return name

        def param(self, name, value):
            return value

        def load_rows(self, name, rows):
            self.tables[name].extend(rows)

        def query(self, sql, params=None):
            run_id = params[0]
            if sql.startswith("DELETE"):
                name = sql.split("`")[1]
                self.tables[name] = [r for r in self.tables[name] if r["run_id"] != run_id]
                return []
            rows = [dict(r, data=b64.b64decode(r["data"]), created_at=datetime.datetime.now()) for r in self.tables[bqmod.FILES] if r["run_id"] == run_id]
            return sorted(rows, key=lambda r: (r["file_id"], r["chunk"]))

    monkeypatch.setattr(bqmod, "CHUNK", 10)
    db = FakeDB()
    files = bqmod.BigQueryFileStore(db)
    store = bqmod.BigQueryStore("p", "d", db)
    data = bytes(range(256)) * 2
    f = StoredFile(id="f1", name="v.mp4", mime="video/mp4", kind="video", size=len(data), data=data)

    async def scenario():
        await files.attach("r1", [f])
        assert len(db.tables[bqmod.FILES]) == -(-len(data) // 10)
        files.forget("r1")
        back = await files.get("r1")
        assert back[0].data == data and back[0].name == "v.mp4"
        store.record_output("r1", {"run_id": "r1", "stage": "analysis"})
        assert await store.delete_run("r1") == "removed"
        assert not any(db.tables.values())

    asyncio.run(scenario())


def test_field_is_detected_from_the_request():
    from app.workflow.profiles import detect_profile

    assert detect_profile("Fix the bug in my Python API endpoint") == "software"
    assert detect_profile("Analyse this sales data and compare the trends", file_names=["sales.csv"]) == "research"
    assert detect_profile("Plan a birthday party for 20 people") == "general"
    with make_client() as client:
        run = client.post("/api/runs", json={"request": "Write a FastAPI backend with a database schema", "require_plan_approval": True}).json()
        assert run["profile"] == "software"
        client.post(f"/api/runs/{run['id']}/cancel")
